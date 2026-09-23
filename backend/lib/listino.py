"""Importazione listino fornitore da file Excel/CSV.

Mappa in modo tollerante le intestazioni italiane più comuni usate dai
fornitori (Gewiss, BTicino, grossisti locali) sui campi del catalogo.
"""

import io
import math
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook

# Sinonimi di intestazione -> campo del modello Materiale.
COLONNE: dict[str, tuple[str, ...]] = {
    "nome": ("nome", "descrizione", "descrizione articolo", "articolo", "denominazione", "prodotto", "desc"),
    "codice_art": ("codice", "codice articolo", "cod articolo", "cod art", "codice art", "sku", "ean", "riferimento"),
    "categoria": ("categoria", "famiglia", "gruppo", "reparto", "tipologia"),
    "unita_misura": ("unita", "unita misura", "unita di misura", "um", "u m", "misura"),
    "prezzo_unitario": ("prezzo", "prezzo vendita", "prezzo unitario", "prezzo pubblico", "prezzo listino", "listino", "vendita"),
    "prezzo_costo": ("costo", "prezzo costo", "prezzo acquisto", "acquisto", "netto", "prezzo netto"),
    "quantita_disponibile": ("quantita", "giacenza", "disponibilita", "disponibile", "scorta", "qta", "q ta", "stock"),
    "scorta_minima": ("scorta minima", "minimo", "soglia", "riordino", "scorta min"),
    "fornitore": ("fornitore", "marca", "produttore", "brand"),
}

UNITA_VALIDE = {"pz", "m", "conf", "rotolo"}
UNITA_SINONIMI = {
    "pezzo": "pz", "pezzi": "pz", "nr": "pz", "n": "pz", "cad": "pz", "pc": "pz", "pz": "pz",
    "metro": "m", "metri": "m", "mt": "m", "ml": "m", "m": "m",
    "confezione": "conf", "conf": "conf", "cf": "conf", "scatola": "conf", "box": "conf",
    "rotolo": "rotolo", "rot": "rotolo", "bobina": "rotolo",
}

CATEGORIE_VALIDE = (
    "Cavi e Conduttori",
    "Magnetotermici e Quadri",
    "Prese e Interruttori",
    "Tubi e Canaline",
    "Illuminazione",
    "Minuteria e Fissaggi",
)

# Parole chiave per indovinare la categoria quando il file non la riporta.
INDIZI_CATEGORIA: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("Cavi e Conduttori", ("cavo", "cavi", "conduttore", "fs17", "fs18", "fg16", "n07v", "unipolare", "trecci")),
    ("Magnetotermici e Quadri", ("magnetotermic", "differenzial", "salvavita", "quadro", "interruttore automatico", "centralino", "modulare")),
    ("Prese e Interruttori", ("presa", "prese", "interruttor", "deviator", "invertitor", "pulsante", "placca", "frutto", "civile")),
    ("Tubi e Canaline", ("tubo", "tubi", "corrugat", "canalina", "canale", "guaina", "passacav", "pressacav", "scatola")),
    ("Illuminazione", ("faretto", "faretti", "led", "lampad", "plafonier", "pannello", "striscia", "applique", "proiettore", "illumin")),
    ("Minuteria e Fissaggi", ("fascett", "morsett", "vite", "viti", "tassell", "capocorda", "collare", "graffett", "nastro", "minuteria")),
)


def _normalizza(testo: object) -> str:
    """Minuscolo, senza accenti e senza punteggiatura: per confrontare intestazioni."""
    s = unicodedata.normalize("NFKD", str(testo)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _mappa_colonne(colonne: list[str]) -> dict[str, str]:
    """Associa ogni colonna del file al campo del catalogo, con match esatto poi parziale."""
    normalizzate = {col: _normalizza(col) for col in colonne}
    mappa: dict[str, str] = {}
    for campo, sinonimi in COLONNE.items():
        for col, norm in normalizzate.items():
            if col in mappa.values() or not norm:
                continue
            if norm in sinonimi:
                mappa[campo] = col
                break
        if campo in mappa:
            continue
        for col, norm in normalizzate.items():
            if col in mappa.values() or not norm:
                continue
            if any(s in norm or norm in s for s in sinonimi):
                mappa[campo] = col
                break
    return mappa


def _numero(valore: object) -> float:
    """Legge un numero in formato italiano o inglese: '1.234,56', '12,5 €', '3.5'."""
    if valore is None:
        return 0.0
    if isinstance(valore, (int, float)):
        return 0.0 if isinstance(valore, float) and math.isnan(valore) else round(float(valore), 2)
    testo = re.sub(r"[^0-9,.\-]", "", str(valore)).strip()
    if not testo:
        return 0.0
    if "," in testo and "." in testo:
        # L'ultimo separatore che appare è quello decimale.
        testo = (
            testo.replace(".", "").replace(",", ".")
            if testo.rfind(",") > testo.rfind(".")
            else testo.replace(",", "")
        )
    elif "," in testo:
        testo = testo.replace(",", ".")
    try:
        return round(float(testo), 2)
    except ValueError:
        return 0.0


def _unita(valore: object) -> str:
    norm = _normalizza(valore).replace(" ", "")
    return UNITA_SINONIMI.get(norm, "pz")


def _categoria(valore: object, nome: str) -> str:
    """Usa la categoria del file se riconoscibile, altrimenti la deduce dal nome."""
    norm = _normalizza(valore)
    if norm:
        for valida in CATEGORIE_VALIDE:
            if _normalizza(valida) == norm:
                return valida
        for valida, indizi in INDIZI_CATEGORIA:
            if any(i in norm for i in indizi):
                return valida
    nome_norm = _normalizza(nome)
    for valida, indizi in INDIZI_CATEGORIA:
        if any(i in nome_norm for i in indizi):
            return valida
    return "Minuteria e Fissaggi"


def leggi_tabella(contenuto: bytes, filename: str) -> pd.DataFrame:
    """Legge xlsx/xls/csv in un DataFrame, saltando eventuali righe di intestazione vuote."""
    nome = (filename or "").lower()
    if nome.endswith(".csv"):
        for sep in (";", ",", "\t"):
            try:
                df = pd.read_csv(io.BytesIO(contenuto), sep=sep, dtype=object)
                if df.shape[1] > 1:
                    return df
            except Exception:
                continue
        return pd.read_csv(io.BytesIO(contenuto), dtype=object)
    df = pd.read_excel(io.BytesIO(contenuto), dtype=object)
    # Alcuni listini hanno titoli/loghi nelle prime righe: cerca la riga di intestazione vera.
    if not _mappa_colonne([str(c) for c in df.columns]).get("nome"):
        for riga in range(min(8, len(df))):
            candidate = [str(v) for v in df.iloc[riga].tolist()]
            if _mappa_colonne(candidate).get("nome"):
                df = pd.read_excel(
                    io.BytesIO(contenuto), dtype=object, header=riga + 1
                )
                break
    return df


def righe_normalizzate(df: pd.DataFrame) -> tuple[list[dict], list[str], dict[str, str]]:
    """Converte il DataFrame in materiali pronti da salvare, più gli avvisi per l'utente."""
    mappa = _mappa_colonne([str(c) for c in df.columns])
    avvisi: list[str] = []
    if "nome" not in mappa:
        return [], ["Nessuna colonna con il nome/descrizione dell'articolo riconosciuta"], mappa
    if "prezzo_unitario" not in mappa:
        avvisi.append("Colonna prezzo non riconosciuta: i prezzi saranno a 0 e potrai correggerli a mano")

    materiali: list[dict] = []
    for idx, riga in df.iterrows():
        nome = str(riga.get(mappa["nome"], "") or "").strip()
        if not nome or _normalizza(nome) in ("nan", "totale", "totali"):
            continue
        prezzo = _numero(riga.get(mappa["prezzo_unitario"])) if "prezzo_unitario" in mappa else 0.0
        costo = _numero(riga.get(mappa["prezzo_costo"])) if "prezzo_costo" in mappa else 0.0
        materiali.append(
            {
                "nome": nome[:160],
                "codice_art": (
                    str(riga.get(mappa["codice_art"], "") or "").strip()[:60]
                    if "codice_art" in mappa
                    else ""
                ),
                "categoria": _categoria(
                    riga.get(mappa["categoria"]) if "categoria" in mappa else "", nome
                ),
                "unita_misura": (
                    _unita(riga.get(mappa["unita_misura"])) if "unita_misura" in mappa else "pz"
                ),
                "prezzo_unitario": prezzo,
                "prezzo_costo": costo,
                "quantita_disponibile": (
                    _numero(riga.get(mappa["quantita_disponibile"]))
                    if "quantita_disponibile" in mappa
                    else 0.0
                ),
                "scorta_minima": (
                    _numero(riga.get(mappa["scorta_minima"])) if "scorta_minima" in mappa else 0.0
                ),
                "fornitore": (
                    str(riga.get(mappa["fornitore"], "") or "").strip()[:80]
                    if "fornitore" in mappa
                    else ""
                ),
            }
        )
    if not materiali:
        avvisi.append("Nessuna riga valida trovata nel file")
    return materiali, avvisi, mappa


def modello_xlsx() -> bytes:
    """File di esempio con le intestazioni attese, da compilare e ricaricare."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Listino"
    intestazioni = [
        "Codice articolo", "Descrizione", "Categoria", "Unità di misura",
        "Prezzo vendita", "Prezzo costo", "Giacenza", "Scorta minima", "Fornitore",
    ]
    ws.append(intestazioni)
    ws.append(["FS18-3G25", "Cavo FS18 3G2,5", "Cavi e Conduttori", "m", "1,20", "0,75", "250", "100", "Gewiss"])
    ws.append(["LN-26602", "Presa bipasso Living Now", "Prese e Interruttori", "pz", "9,80", "5,90", "24", "10", "BTicino"])
    ws.append(["FAR-LED7", "Faretto LED 7W orientabile", "Illuminazione", "pz", "14,50", "8,90", "18", "8", "Illuminart"])
    for i, testo in enumerate(intestazioni, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(testo) + 4)
        ws.cell(row=1, column=i).font = ws.cell(row=1, column=i).font.copy(bold=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
